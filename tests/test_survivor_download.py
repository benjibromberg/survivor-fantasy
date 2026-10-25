"""Tests for download_survivor_data(): how survivoR.xlsx reaches the disk.

Two gunicorn workers each run the daily refresh, so both download at the same
moment into the same file, and request handlers read it at any time. Writing
the file in place let one of them read a half-written workbook ("Truncated
file header"). It is now written to a temporary file in the same directory and
swapped in with os.replace, so the path only ever holds a complete file.

The file is built from survivoR's dev/json tables (tests/test_survivor_json.py
covers how), all fetched at the one commit GitHub reports as survivoR's head,
and nothing is written until every table has arrived, so a failed fetch leaves
yesterday's file in place.
"""

import datetime
import json
import os
import pathlib
import subprocess
import sys
import tempfile
import time
from types import SimpleNamespace

import pandas as pd
import pytest
import requests

OLD_BYTES = b"the previous survivoR.xlsx"


class _FakeResponse:
    def __init__(self, status, body=b"", headers=None):
        self.status_code = status
        self.content = body
        self.headers = headers or {}

    def raise_for_status(self):
        if self.status_code >= 400:
            raise requests.HTTPError(f"{self.status_code} from upstream")

    def json(self):
        return json.loads(self.content)


@pytest.fixture()
def target(tmp_path, monkeypatch):
    """SURVIVOR_DATA_FILE in an empty temp directory."""
    path = tmp_path / "survivoR.xlsx"
    monkeypatch.setattr("app.data.SURVIVOR_DATA_FILE", str(path))
    return path


def _leftovers(path):
    """Anything in the data directory other than the file itself."""
    return sorted(p.name for p in path.parent.iterdir() if p != path)


def _writes(body):
    """A write callback that puts `body` at the path it is given."""

    def write(path):
        with open(path, "wb") as f:
            f.write(body)

    return write


class TestWriteAtomically:
    def test_a_successful_write_replaces_the_file(self, target):
        from app.data import _write_atomically

        target.write_bytes(OLD_BYTES)

        _write_atomically(str(target), _writes(b"new"))

        assert target.read_bytes() == b"new"
        assert _leftovers(target) == []

    def test_a_write_that_fails_partway_keeps_the_previous_file(self, target):
        from app.data import _write_atomically

        target.write_bytes(OLD_BYTES)

        def half_then_fail(tmp):
            with open(tmp, "wb") as f:
                f.write(b"half a work")
            raise OSError("disk full")

        with pytest.raises(OSError, match="disk full"):
            _write_atomically(str(target), half_then_fail)

        assert target.read_bytes() == OLD_BYTES
        assert _leftovers(target) == []

    def test_a_failed_first_write_creates_no_file(self, target):
        from app.data import _write_atomically

        def fail(tmp):
            with open(tmp, "wb") as f:
                f.write(b"half")
            raise OSError("disk full")

        with pytest.raises(OSError):
            _write_atomically(str(target), fail)

        assert not target.exists()
        assert _leftovers(target) == []

    def test_the_file_arrives_whole_by_rename_from_the_same_directory(
        self, target, monkeypatch
    ):
        """A rename within one directory is atomic; a copy or a write is not."""
        from app import data

        swaps = []
        real_replace = os.replace

        def spy(src, dst):
            with open(src, "rb") as f:
                swaps.append((os.path.dirname(src), f.read(), dst))
            real_replace(src, dst)

        monkeypatch.setattr(data.os, "replace", spy)

        data._write_atomically(str(target), _writes(b"whole"))

        assert swaps == [(str(target.parent), b"whole", str(target))]

    def test_two_writers_never_share_a_temporary_file(self, target):
        """Both workers download at 08:00; each must write its own temp file."""
        from app.data import _write_atomically

        seen = []

        def inner(tmp):
            seen.append(tmp)
            _writes(b"inner")(tmp)

        def outer(tmp):
            seen.append(tmp)
            _write_atomically(str(target), inner)
            _writes(b"outer")(tmp)

        _write_atomically(str(target), outer)

        assert len(set(seen)) == 2
        assert target.read_bytes() == b"outer"
        assert _leftovers(target) == []


# The fixture tables are survivoR's own dev/json files, trimmed (see
# tests/test_survivor_json.py for what they hold).
FIXTURE_JSON = (
    pathlib.Path(__file__).resolve().parent / "fixtures" / "survivoR" / "json-c720d36"
)
FIXTURE_COMMIT = "c720d36c6c8706c052a822bcda850ee273d08c49"
COMMIT_URL = "https://api.github.com/repos/doehm/survivoR/commits/master"
RAW_BASE = "https://raw.githubusercontent.com/doehm/survivoR/"


def _url(table, commit=FIXTURE_COMMIT):
    return f"{RAW_BASE}{commit}/dev/json/{table}.json"


@pytest.fixture()
def upstream(monkeypatch):
    """Serve survivoR from the fixtures, by URL: GitHub's answer for the head
    commit of survivoR's master, and the dev/json tables at that commit.

    `served[url]` can be replaced with an int (an HTTP status), a (status,
    headers) pair, an exception to raise, other bytes, or a list of those
    answered one per request, the last again once the rest are used. `calls` lists the URLs requested, in
    order, `headers` the headers sent with each, and `waits` the seconds the
    download asked to sleep before a retry (it never actually sleeps).
    """
    from app.data import SURVIVOR_TABLES

    served = {COMMIT_URL: FIXTURE_COMMIT.encode()}
    for table in SURVIVOR_TABLES.values():
        served[_url(table)] = (FIXTURE_JSON / f"{table}.json").read_bytes()
    calls = []
    sent = []
    waits = []

    def fake_get(url, timeout=None, headers=None, **_kwargs):
        calls.append(url)
        sent.append(headers or {})
        body = served.get(url, 404)
        if isinstance(body, list):
            body = body.pop(0) if len(body) > 1 else body[0]
        if isinstance(body, Exception):
            raise body
        if isinstance(body, int):
            return _FakeResponse(body)
        if isinstance(body, tuple):
            return _FakeResponse(body[0], headers=body[1])
        return _FakeResponse(200, body)

    monkeypatch.setattr("app.data.requests.get", fake_get)
    monkeypatch.setattr("app.data.time.sleep", waits.append)
    return SimpleNamespace(served=served, calls=calls, headers=sent, waits=waits)


def _sheet_names(path):
    with pd.ExcelFile(path) as book:
        return sorted(book.sheet_names)


class TestDownloadSurvivorData:
    def test_fetches_each_table_the_app_reads_once(self, target, upstream):
        from app import data

        data.download_survivor_data()

        assert upstream.calls == [COMMIT_URL] + [
            _url(table) for table in data.SURVIVOR_TABLES.values()
        ]

    def test_writes_every_sheet_into_survivor_xlsx(self, target, upstream):
        from app import data

        data.download_survivor_data()

        assert _sheet_names(target) == sorted(data.SURVIVOR_TABLES)
        castaways = pd.read_excel(target, "Castaways")
        assert set(castaways["season"]) == {50, 51}

    def test_a_download_replaces_the_previous_file(self, target, upstream):
        from app import data

        target.write_bytes(OLD_BYTES)

        data.download_survivor_data()

        assert target.read_bytes() != OLD_BYTES
        assert "Castaways" in _sheet_names(target)
        assert _leftovers(target) == []

    @pytest.mark.parametrize(
        "failure",
        [404, 503, requests.ConnectionError("reset"), requests.Timeout("slow")],
        ids=["404", "503", "connection-error", "timeout"],
    )
    def test_a_failed_fetch_keeps_the_previous_file(self, target, upstream, failure):
        from app import data

        target.write_bytes(OLD_BYTES)
        upstream.served[_url("tribe_mapping")] = failure

        with pytest.raises(requests.RequestException):
            data.download_survivor_data()

        assert target.read_bytes() == OLD_BYTES
        assert _leftovers(target) == []

    @pytest.mark.parametrize(
        "body",
        [b"[]", b"{}", b'["not a row"]', b"<html>rate limited</html>"],
        ids=["empty-list", "object", "list-of-strings", "not-json"],
    )
    def test_a_table_that_is_not_rows_keeps_the_previous_file(
        self, target, upstream, body
    ):
        from app import data

        target.write_bytes(OLD_BYTES)
        upstream.served[_url("castaways")] = body

        with pytest.raises(ValueError):
            data.download_survivor_data()

        assert target.read_bytes() == OLD_BYTES
        assert _leftovers(target) == []

    @pytest.mark.parametrize(
        ("failure", "error"),
        [(503, requests.HTTPError), (b"{}", ValueError)],
        ids=["fetch-fails", "not-rows"],
    )
    def test_nothing_is_written_until_every_table_has_arrived(
        self, target, upstream, monkeypatch, failure, error
    ):
        """A failure on the last table must not have started a write."""
        from app import data

        writes = []
        monkeypatch.setattr(data, "_write_atomically", lambda *a: writes.append(a))
        last = list(data.SURVIVOR_TABLES.values())[-1]
        upstream.served[_url(last)] = failure

        with pytest.raises(error):
            data.download_survivor_data()

        assert writes == []

    def test_a_failed_first_download_creates_no_file(self, target, upstream):
        from app import data

        upstream.served[_url("episodes")] = 503

        with pytest.raises(requests.HTTPError):
            data.download_survivor_data()

        assert not target.exists()
        assert _leftovers(target) == []

    def test_the_download_goes_through_the_atomic_swap(
        self, target, upstream, monkeypatch
    ):
        from app import data

        swaps = []
        real_replace = os.replace

        def spy(src, dst):
            assert os.path.dirname(src) == os.path.dirname(dst)
            assert _sheet_names(src) == sorted(data.SURVIVOR_TABLES)
            swaps.append(dst)
            real_replace(src, dst)

        monkeypatch.setattr(data.os, "replace", spy)

        data.download_survivor_data()

        assert swaps == [str(target)]


class TestOneCommitPerRefresh:
    """raw.githubusercontent.com caches each file for five minutes, so for a
    while after an upstream push the tables fetched from master could come
    from two commits. The head commit is now resolved once, through GitHub's
    API, and every table is fetched at that commit."""

    def test_every_table_is_fetched_at_the_commit_resolved_first(
        self, target, upstream
    ):
        from app import data

        data.download_survivor_data()

        assert upstream.calls[0] == COMMIT_URL
        assert upstream.calls[1:] == [
            _url(table) for table in data.SURVIVOR_TABLES.values()
        ]

    def test_the_commit_is_asked_for_as_a_bare_sha(self, target, upstream):
        """GitHub answers with just the 40 characters for this media type."""
        from app import data

        data.download_survivor_data()

        assert upstream.headers[0].get("Accept") == "application/vnd.github.sha"

    def test_returns_the_commit_it_built_from(self, target, upstream):
        from app import data

        upstream.served[COMMIT_URL] = f"{FIXTURE_COMMIT}\n".encode()

        assert data.download_survivor_data() == FIXTURE_COMMIT

    def test_a_push_during_the_refresh_does_not_mix_commits(self, target, upstream):
        """Upstream pushes while the tables arrive: master serves the new
        commit from then on, and GitHub would report it as the head."""
        from app import data

        newer = "0123456789abcdef0123456789abcdef01234567"
        upstream.served[COMMIT_URL] = [FIXTURE_COMMIT.encode(), newer.encode()]
        for table in data.SURVIVOR_TABLES.values():
            upstream.served[_url(table, "master")] = b'[{"season": 99}]'
            upstream.served[_url(table, newer)] = b'[{"season": 99}]'

        data.download_survivor_data()

        assert upstream.calls.count(COMMIT_URL) == 1
        assert [url for url in upstream.calls[1:] if FIXTURE_COMMIT not in url] == []
        assert set(pd.read_excel(target, "Castaways")["season"]) == {50, 51}


class TestACommitLookupThatFails:
    """The commit lookup is GitHub's API, which allows an address 60 requests
    an hour without a token, shared by everything behind that address; past
    that it answers 403. The fetch from master never needed it, so a failed
    lookup must not cost the refresh: the tables come from master as before,
    and the log says this refresh is not pinned to one commit."""

    @pytest.fixture()
    def master(self, upstream):
        """Serve the fixture tables at master too."""
        from app.data import SURVIVOR_TABLES

        for table in SURVIVOR_TABLES.values():
            upstream.served[_url(table, "master")] = upstream.served[_url(table)]
        return upstream

    @pytest.mark.parametrize(
        "failure",
        [
            403,
            404,
            429,
            500,
            503,
            requests.ConnectionError("reset"),
            requests.Timeout("slow"),
        ],
        ids=[
            "403-allowance-used",
            "404",
            "429",
            "500",
            "503",
            "connection-error",
            "timeout",
        ],
    )
    def test_the_tables_come_from_master_instead(self, target, master, caplog, failure):
        from app import data

        target.write_bytes(OLD_BYTES)
        master.served[COMMIT_URL] = failure

        with caplog.at_level("WARNING", logger="app.data"):
            assert data.download_survivor_data() == data.SURVIVOR_BRANCH

        tables = [url for url in master.calls if url != COMMIT_URL]
        assert tables == [_url(t, "master") for t in data.SURVIVOR_TABLES.values()]
        assert "Castaways" in _sheet_names(target)
        assert any("not pinned to one commit" in r.getMessage() for r in caplog.records)

    @pytest.mark.parametrize(
        "body",
        [b"<html>try later</html>", b"", b"c720d36", b"master", b"{}"],
        ids=["html", "empty", "short-sha", "branch-name", "json"],
    )
    def test_a_reply_that_is_not_a_commit_is_never_put_in_a_url(
        self, target, master, body
    ):
        from app import data

        master.served[COMMIT_URL] = body

        assert data.download_survivor_data() == data.SURVIVOR_BRANCH

        tables = [url for url in master.calls if url != COMMIT_URL]
        assert tables == [_url(t, "master") for t in data.SURVIVOR_TABLES.values()]

    def test_an_allowance_used_up_is_not_retried(self, target, master):
        """GitHub's 403 lasts until the hour resets; asking again only spends
        another request."""
        from app import data

        master.served[COMMIT_URL] = 403

        data.download_survivor_data()

        assert master.calls.count(COMMIT_URL) == 1
        assert master.waits == []

    def test_a_failure_after_the_fallback_still_keeps_the_previous_file(
        self, target, master
    ):
        from app import data

        target.write_bytes(OLD_BYTES)
        master.served[COMMIT_URL] = 403
        master.served[_url("episodes", "master")] = 404

        with pytest.raises(requests.HTTPError):
            data.download_survivor_data()

        assert target.read_bytes() == OLD_BYTES
        assert _leftovers(target) == []

    def test_a_working_lookup_still_pins_every_table(self, target, master):
        """The fallback must not become the normal path."""
        from app import data

        assert data.download_survivor_data() == FIXTURE_COMMIT
        assert [url for url in master.calls if "/master/" in url] == []


TRANSIENT = [
    500,
    502,
    503,
    504,
    429,
    requests.ConnectionError("reset"),
    requests.Timeout("slow"),
    requests.exceptions.ChunkedEncodingError("cut off"),
]
TRANSIENT_IDS = [
    "500",
    "502",
    "503",
    "504",
    "429",
    "connection-error",
    "timeout",
    "cut-off",
]


class TestATransientFailureIsRetried:
    """A 5xx or a timeout on any one of the requests used to skip the whole
    refresh. A failure likely to pass is now retried, waiting longer before
    each retry; one that will not pass (a 4xx) is not."""

    @pytest.mark.parametrize("failure", TRANSIENT, ids=TRANSIENT_IDS)
    def test_a_table_that_fails_once_is_fetched_again(self, target, upstream, failure):
        from app import data

        url = _url("castaways")
        upstream.served[url] = [failure, upstream.served[url]]

        data.download_survivor_data()

        assert upstream.calls.count(url) == 2
        assert set(pd.read_excel(target, "Castaways")["season"]) == {50, 51}
        assert len(upstream.waits) == 1

    def test_the_commit_lookup_is_retried_too(self, target, upstream):
        from app import data

        upstream.served[COMMIT_URL] = [503, FIXTURE_COMMIT.encode()]

        assert data.download_survivor_data() == FIXTURE_COMMIT
        assert upstream.calls[:2] == [COMMIT_URL, COMMIT_URL]

    def test_each_retry_waits_longer_than_the_last(self, target, upstream):
        from app import data

        url = _url("episodes")
        upstream.served[url] = [503] * (data.FETCH_ATTEMPTS - 1) + [
            upstream.served[url]
        ]

        data.download_survivor_data()

        waits = upstream.waits
        assert len(waits) == data.FETCH_ATTEMPTS - 1 >= 2
        assert waits[0] > 0
        assert waits == sorted(set(waits))

    @pytest.mark.parametrize("failure", TRANSIENT, ids=TRANSIENT_IDS)
    def test_a_failure_that_persists_gives_up_and_keeps_the_previous_file(
        self, target, upstream, failure
    ):
        from app import data

        target.write_bytes(OLD_BYTES)
        url = _url("vote_history")
        upstream.served[url] = failure

        with pytest.raises(requests.RequestException):
            data.download_survivor_data()

        assert upstream.calls.count(url) == data.FETCH_ATTEMPTS
        assert len(upstream.waits) == data.FETCH_ATTEMPTS - 1
        assert target.read_bytes() == OLD_BYTES
        assert _leftovers(target) == []

    @pytest.mark.parametrize("status", [400, 401, 403, 404])
    def test_a_client_error_is_not_retried(self, target, upstream, status):
        from app import data

        url = _url("jury_votes")
        upstream.served[url] = status

        with pytest.raises(requests.HTTPError):
            data.download_survivor_data()

        assert upstream.calls.count(url) == 1
        assert upstream.waits == []

    def test_a_retry_after_header_is_waited_out(self, target, upstream):
        from app import data

        url = _url("castaways")
        upstream.served[url] = [(429, {"Retry-After": "7"}), upstream.served[url]]

        data.download_survivor_data()

        assert upstream.waits == [7]
        assert upstream.calls.count(url) == 2

    def test_a_retry_after_shorter_than_the_backoff_keeps_the_backoff(
        self, target, upstream
    ):
        from app import data

        url = _url("castaways")
        upstream.served[url] = [(503, {"Retry-After": "0"}), upstream.served[url]]

        data.download_survivor_data()

        assert upstream.waits == [data.FETCH_RETRY_WAIT]

    def test_a_retry_after_given_as_a_date_is_waited_out(self, target, upstream):
        from email.utils import format_datetime

        from app import data

        url = _url("castaways")
        when = format_datetime(
            datetime.datetime.now(datetime.UTC) + datetime.timedelta(seconds=12),
            usegmt=True,
        )
        upstream.served[url] = [(429, {"Retry-After": when}), upstream.served[url]]

        data.download_survivor_data()

        assert len(upstream.waits) == 1
        assert 10 <= upstream.waits[0] <= 13

    def test_a_retry_after_too_long_to_wait_fails_at_once(self, target, upstream):
        """An hour's wait would hold the refresh for an hour; give up instead."""
        from app import data

        target.write_bytes(OLD_BYTES)
        url = _url("castaways")
        upstream.served[url] = (429, {"Retry-After": "3600"})

        with pytest.raises(requests.HTTPError):
            data.download_survivor_data()

        assert upstream.calls.count(url) == 1
        assert upstream.waits == []
        assert target.read_bytes() == OLD_BYTES

    def test_a_table_that_is_not_rows_is_not_retried(self, target, upstream):
        """A bad body is upstream's data, not a failed request."""
        from app import data

        url = _url("castaways")
        upstream.served[url] = b"{}"

        with pytest.raises(ValueError):
            data.download_survivor_data()

        assert upstream.calls.count(url) == 1
        assert upstream.waits == []


def _dead_pid():
    """The pid of a process that has already exited."""
    done = subprocess.run(
        [sys.executable, "-c", "import os; print(os.getpid())"],
        capture_output=True,
        text=True,
        check=True,
    )
    return int(done.stdout)


def _temp_beside(target, pid):
    """A temp file named the way _write_atomically names its own."""
    path = target.parent / f"{target.name}.{pid}.0badf00d.tmp"
    path.write_bytes(b"half a workbook")
    return path


class TestTempFilesAKilledWriterLeft:
    """_write_atomically removes its temp file when the write raises, but a
    worker killed outright (SIGKILL after a timeout, out of memory, a
    container stop) runs no cleanup, and its survivoR.xlsx.<pid>.<random>.tmp
    stayed in the data directory for good. The next build removes it."""

    def test_a_killed_writers_temp_file_goes_with_the_next_build(
        self, target, upstream
    ):
        from app import data

        writer = (
            "import os, sys, time\n"
            "path = f'{sys.argv[1]}.{os.getpid()}.0badf00d.tmp'\n"
            "open(path, 'wb').write(b'half a workbook')\n"
            "print('writing', flush=True)\n"
            "time.sleep(60)\n"
        )
        proc = subprocess.Popen(
            [sys.executable, "-c", writer, str(target)],
            stdout=subprocess.PIPE,
            text=True,
        )
        try:
            assert proc.stdout.readline().strip() == "writing"
        finally:
            proc.kill()
            proc.wait()
            proc.stdout.close()
        assert _leftovers(target) == [f"{target.name}.{proc.pid}.0badf00d.tmp"]

        data.download_survivor_data()

        assert _leftovers(target) == []

    def test_a_build_that_fails_removes_a_dead_writers_file_too(self, target, upstream):
        """The sweep comes first, so a day whose build fails still cleans up."""
        from app import data

        _temp_beside(target, _dead_pid())
        upstream.served[_url("castaways")] = 404

        with pytest.raises(requests.HTTPError):
            data.download_survivor_data()

        assert _leftovers(target) == []

    def test_a_running_writers_temp_file_is_left_alone(self, target, upstream):
        """Both workers build at 08:00: the other one's file is mid-write."""
        from app import data

        live = _temp_beside(target, os.getppid())

        data.download_survivor_data()

        assert _leftovers(target) == [live.name]

    def test_an_hour_old_temp_file_goes_whatever_its_pid(self, target, upstream):
        """A restarted container reuses pids; no build takes an hour."""
        from app import data

        old = _temp_beside(target, os.getppid())
        an_hour_and_a_bit_ago = time.time() - 3700
        os.utime(old, (an_hour_and_a_bit_ago, an_hour_and_a_bit_ago))

        data.download_survivor_data()

        assert _leftovers(target) == []

    def test_nothing_else_in_the_data_directory_is_touched(self, target, upstream):
        from app import data

        dead = _dead_pid()
        others = [
            "survivor_fantasy.db",
            "survivoR.xlsx.bak",
            f"survivoR.xlsx.{dead}.tmp",
            f"other.xlsx.{dead}.0badf00d.tmp",
            f"survivoR.xlsx.{dead}.0badf00d.tmp.keep",
        ]
        for name in others:
            (target.parent / name).write_bytes(b"not ours")

        data.download_survivor_data()

        assert _leftovers(target) == sorted(others)


@pytest.fixture()
def system_temp(tmp_path_factory, monkeypatch):
    """An empty directory standing in for the system temp directory.

    openpyxl's write-only mode streams each sheet into a file of its own there
    (openpyxl.*) and zips them into the workbook on save. It deletes them on
    save, or at process exit, and a gunicorn worker runs for weeks.
    """
    path = tmp_path_factory.mktemp("system-temp")
    monkeypatch.setattr(tempfile, "tempdir", str(path))
    return path


def _with_value(upstream, table, column, value):
    """Serve `table` with `column` of its first row set to `value`."""
    records = json.loads(upstream.served[_url(table)])
    records[0][column] = value
    upstream.served[_url(table)] = json.dumps(records).encode()


# Values that pass the shape check but that the workbook cannot hold. The
# first is a date column (every other episode_date is YYYY-MM-DD), so the
# build turns it into a date, and there is no 30 February.
UNWRITABLE = [
    ("episodes", "episode_date", "2026-02-30"),
    ("castaways", "castaway", "Q\x07"),
    ("castaways", "castaway", ["Q"]),
    ("castaways", "castaway", {"name": "Q"}),
]
UNWRITABLE_IDS = ["impossible-date", "control-character", "list", "object"]


class TestAValueTheWorkbookCannotHold:
    """One bad value used to pass the check, then abort the write halfway and
    leave openpyxl's per-sheet temp files behind, every day until upstream
    fixed it. It is now refused with the tables, before any writing starts."""

    @pytest.mark.parametrize(
        ("table", "column", "value"), UNWRITABLE, ids=UNWRITABLE_IDS
    )
    def test_is_refused_before_the_workbook_is_started(
        self, target, upstream, monkeypatch, table, column, value
    ):
        from app import data

        writes = []
        monkeypatch.setattr(
            data, "write_survivor_workbook", lambda *a: writes.append(a)
        )
        _with_value(upstream, table, column, value)

        with pytest.raises(ValueError, match=f"{table}.*{column}"):
            data.download_survivor_data()

        assert writes == []

    @pytest.mark.parametrize(
        ("table", "column", "value"), UNWRITABLE, ids=UNWRITABLE_IDS
    )
    def test_keeps_the_previous_file_and_leaves_no_temp_files(
        self, target, upstream, system_temp, table, column, value
    ):
        from app import data

        target.write_bytes(OLD_BYTES)
        _with_value(upstream, table, column, value)

        with pytest.raises(ValueError):
            data.download_survivor_data()

        assert target.read_bytes() == OLD_BYTES
        assert _leftovers(target) == []
        assert sorted(p.name for p in system_temp.iterdir()) == []

    def test_the_tables_as_published_pass(self, target, upstream, system_temp):
        """The check must not refuse the real data."""
        from app import data

        data.download_survivor_data()

        assert _sheet_names(target) == sorted(data.SURVIVOR_TABLES)
        assert sorted(p.name for p in system_temp.iterdir()) == []


class TestAWriteThatFailsPartway:
    """Whatever the check above does not foresee, a failed write still
    removes openpyxl's temp files instead of leaving them to process exit."""

    def test_a_failure_between_sheets_leaves_no_temp_files(self, tmp_path, system_temp):
        from app.data import write_survivor_workbook

        def tables():
            yield "First", [{"season": 51, "castaway": "Q"}] * 50
            yield "Second", [{"season": 51}]
            raise OSError("disk full")

        with pytest.raises(OSError, match="disk full"):
            write_survivor_workbook(str(tmp_path / "book.xlsx"), tables())

        assert sorted(p.name for p in system_temp.iterdir()) == []

    def test_a_cell_openpyxl_refuses_leaves_no_temp_files(self, tmp_path, system_temp):
        """Straight to the writer, past the check: openpyxl refuses the cell."""
        from openpyxl.utils.exceptions import IllegalCharacterError

        from app.data import write_survivor_workbook

        tables = [("First", [{"castaway": "Q"}]), ("Second", [{"castaway": "Q\x07"}])]

        with pytest.raises(IllegalCharacterError):
            write_survivor_workbook(str(tmp_path / "book.xlsx"), tables)

        assert sorted(p.name for p in system_temp.iterdir()) == []

    def test_the_guard_sees_openpyxl_temp_files(self, tmp_path, system_temp):
        """The temp directory above is the one openpyxl writes to, so an empty
        one means the files were removed, not written somewhere else."""
        from openpyxl import Workbook

        workbook = Workbook(write_only=True)
        workbook.create_sheet("Only").append(["season"])

        try:
            assert [p.name.startswith("openpyxl.") for p in system_temp.iterdir()] == [
                True
            ]
        finally:
            workbook.save(str(tmp_path / "book.xlsx"))

        assert sorted(p.name for p in system_temp.iterdir()) == []


class TestTextStartingWithEquals:
    """openpyxl takes a string starting with = for a formula. A formula cell
    has no stored value, so pd.read_excel read a castaway named =Q back as
    NaN. survivoR's values are data, never formulas."""

    def test_a_castaway_named_like_a_formula_reads_back_as_text(self, target, upstream):
        from app import data

        _with_value(upstream, "castaways", "castaway", "=Q")

        data.download_survivor_data()

        castaways = pd.read_excel(target, "Castaways")
        assert castaways["castaway"].iloc[0] == "=Q"

    def test_no_cell_is_written_as_a_formula(self, tmp_path):
        """Column names too, and whatever follows the =."""
        from openpyxl import load_workbook

        from app.data import write_survivor_workbook

        path = tmp_path / "book.xlsx"
        records = [{"=name": "=Q", "season": 51, "note": "=SUM(A1:A2)", "eq": "="}]

        write_survivor_workbook(str(path), [("Sheet", records)])

        frame = pd.read_excel(path, "Sheet")
        assert list(frame.columns) == ["=name", "season", "note", "eq"]
        assert frame.iloc[0].tolist() == ["=Q", 51, "=SUM(A1:A2)", "="]
        book = load_workbook(path)
        try:
            types = {
                c.value: c.data_type for row in book["Sheet"].iter_rows() for c in row
            }
        finally:
            book.close()
        assert types["=Q"] == "s"
        assert types["=SUM(A1:A2)"] == "s"
        assert types["=name"] == "s"
        assert types[51] == "n"


class TestOneWayIn:
    """Every copy of survivoR.xlsx comes from download_survivor_data()."""

    def _python_sources(self):
        repo = pathlib.Path(__file__).resolve().parent.parent
        return sorted((repo / "app").rglob("*.py")) + sorted(repo.glob("*.py"))

    def test_the_scan_reads_the_scripts_too(self):
        names = {path.name for path in self._python_sources()}

        assert {"data.py", "seed.py", "add_season.py", "analyze_scoring.py"} <= names

    def test_nothing_fetches_survivors_xlsx_export(self):
        """Upstream regenerates the xlsx only now and then; it went stale."""
        stale = [
            path.name
            for path in self._python_sources()
            if "dev/xlsx" in path.read_text()
        ]

        assert stale == []

    def test_the_json_source_is_still_named(self):
        """The guard above must not pass because the download vanished."""
        sources = {path.name: path.read_text() for path in self._python_sources()}

        assert "/dev/json/" in sources["data.py"]
