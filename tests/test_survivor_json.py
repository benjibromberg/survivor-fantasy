"""survivoR.xlsx is built from survivoR's dev/json tables, not downloaded.

survivoR regenerates dev/json/<table>.json with every episode but its xlsx
export only now and then: the xlsx stopped at US51 episode 2 (commit c720d36,
"ADD US51E02") while episodes 3 and 4 reached dev/json. The app read the xlsx,
so it showed a stale season. download_survivor_data() now fetches each table
the app reads from dev/json and writes the same sheets into survivoR.xlsx, so
nothing that reads the workbook changes.

The fixtures in tests/fixtures/survivoR/ are both exports at survivoR commit
c720d36, which regenerated the xlsx and the JSON together, so they hold the
same data in two formats. Each keeps US seasons 50 and 51 (51 has episodes 1
and 2 at that commit) of every table the app reads; Castaway Details has no
season column and keeps those seasons' castaways. To stay small, the five
per-episode tables (Boot Mapping, Challenge Results, Confessionals, Tribe
Mapping, Vote History) keep only season 50's episodes 1, 2, 12 and 13. The
xlsx was trimmed cell by cell, keeping each cell's upstream type and format.
"""

import datetime as dt
import importlib
import json
import pathlib
import re
import sys

import pandas as pd
import pytest

REPO = pathlib.Path(__file__).resolve().parent.parent
FIXTURES = REPO / "tests" / "fixtures" / "survivoR"
FIXTURE_JSON = FIXTURES / "json-c720d36"
FIXTURE_XLSX = FIXTURES / "upstream-c720d36.xlsx"

# survivoR's dev/json carries at most 4 decimal places (0.803225 in the xlsx is
# 0.8032 in the JSON). Floats are compared to that precision; nothing else is.
JSON_DECIMALS_TOLERANCE = 5e-5


def _fixture_tables():
    """{sheet: records} for every table the app reads, from the fixture JSON."""
    from app.data import SURVIVOR_TABLES

    return {
        sheet: json.loads((FIXTURE_JSON / f"{table}.json").read_text())
        for sheet, table in SURVIVOR_TABLES.items()
    }


@pytest.fixture(scope="module")
def built_xlsx(tmp_path_factory):
    """The workbook download_survivor_data() would write from the fixture JSON."""
    from app.data import write_survivor_workbook

    path = tmp_path_factory.mktemp("built") / "survivoR.xlsx"
    write_survivor_workbook(str(path), _fixture_tables().items())
    return path


# ── Which sheets are built ─────────────────────────────────────────────────

# A read_excel() call and the sheet name it reads, across a line break too.
READ_EXCEL_SHEET = re.compile(
    r"read_excel\(\s*[^,()]+,\s*(?:sheet_name\s*=\s*)?[\"']([^\"']+)[\"']"
)


def _sheets_read_by_the_code():
    """{sheet name: [files]} for every read_excel() in the app and scripts."""
    sources = sorted((REPO / "app").rglob("*.py")) + sorted(REPO.glob("*.py"))
    found = {}
    for path in sources:
        text = path.read_text()
        sheets = READ_EXCEL_SHEET.findall(text)
        # A read whose sheet name is not a string literal would be missed;
        # fail rather than let it through.
        assert len(sheets) == text.count("read_excel("), (
            f"{path.name} has a read_excel() call this scan cannot read"
        )
        for sheet in sheets:
            found.setdefault(sheet, []).append(path.relative_to(REPO).as_posix())
    return found


class TestWhichSheetsAreBuilt:
    def test_the_scan_finds_the_reads_it_is_meant_to(self):
        """Guard the scan itself, so the two tests below cannot pass vacuously."""
        found = _sheets_read_by_the_code()

        assert "app/data.py" in found["Castaways"]
        assert "app/predictions.py" in found["Advantage Movement"]
        assert "seed.py" in found["Castaway Details"]
        assert "analyze_scoring.py" in found["Vote History"]

    def test_every_sheet_the_code_reads_is_built(self):
        from app.data import SURVIVOR_TABLES

        missing = set(_sheets_read_by_the_code()) - set(SURVIVOR_TABLES)

        assert not missing, f"read but never built into survivoR.xlsx: {missing}"

    def test_every_sheet_built_is_read_somewhere(self):
        """No table is fetched every day for nothing."""
        from app.data import SURVIVOR_TABLES

        unread = set(SURVIVOR_TABLES) - set(_sheets_read_by_the_code())

        assert not unread, f"built into survivoR.xlsx but never read: {unread}"

    def test_each_sheet_comes_from_its_dev_json_table(self):
        from app.data import SURVIVOR_TABLES

        for sheet, table in SURVIVOR_TABLES.items():
            assert table == sheet.lower().replace(" ", "_")
            assert (FIXTURE_JSON / f"{table}.json").exists()


# ── Turning one JSON table into worksheet rows ─────────────────────────────


def _rows(records):
    from app.data import survivor_sheet_rows

    header, *rows = survivor_sheet_rows(records)
    return header, rows


class TestSurvivorSheetRows:
    def test_values_keep_their_json_types(self):
        header, rows = _rows(
            [{"season": 51, "castaway": "Q", "jury": False, "score": 0.8032}]
        )

        assert header == ["season", "castaway", "jury", "score"]
        assert rows == [[51, "Q", False, 0.8032]]
        assert type(rows[0][2]) is bool

    def test_a_key_left_out_of_a_row_is_an_empty_cell(self):
        """survivoR's JSON omits a null field from the row rather than writing null."""
        _, rows = _rows([{"order": 1, "place": 21}, {"order": 2}])

        assert rows == [[1, 21], [2, None]]

    def test_a_column_missing_from_the_first_row_is_still_a_column(self):
        header, rows = _rows([{"version": "US"}, {"version": "US", "n_jury": 8}])

        assert header == ["version", "n_jury"]
        assert rows == [["US", None], ["US", 8]]

    def test_a_column_of_iso_dates_becomes_dates(self):
        """The xlsx stores R Dates as Excel dates; the JSON as YYYY-MM-DD text."""
        _, rows = _rows([{"episode_date": "2026-10-15"}, {}])

        assert rows == [[dt.datetime(2026, 10, 15)], [None]]

    def test_a_column_with_any_other_text_stays_text(self):
        _, rows = _rows([{"premiered": "2026-09-23"}, {"premiered": "TBA"}])

        assert rows == [["2026-09-23"], ["TBA"]]

    def test_a_date_with_a_time_stays_text(self):
        _, rows = _rows([{"aired": "2026-09-23 20:00:00"}])

        assert rows == [["2026-09-23 20:00:00"]]

    def test_numbers_are_never_taken_for_dates(self):
        _, rows = _rows([{"season": 20260923}])

        assert rows == [[20260923]]

    def test_a_column_with_no_value_in_any_row_has_nothing_to_write(self):
        """The one thing the JSON cannot carry: a column that is null in every
        row leaves no key behind, so the sheet has no such column, where the
        xlsx has an empty one. None of survivoR's full tables has such a
        column; the trimmed fixtures do, and are compared with that in mind."""
        header, rows = _rows([{"version": "US"}, {"version": "US"}])

        assert header == ["version"]
        assert rows == [["US"], ["US"]]


class TestWriteSurvivorWorkbook:
    def test_each_sheet_states_its_size_up_front(self, built_xlsx):
        """Opening a workbook, openpyxl reads each sheet's size from the
        <dimension> element at its top, and parses the whole sheet when it is
        missing. Write-only mode leaves it out unless told the size, and then
        every pd.read_excel call on the full workbook spent over 2 seconds
        scanning all 14 sheets, against a tenth of a second for upstream's."""
        from openpyxl import load_workbook

        book = load_workbook(built_xlsx, read_only=True)
        try:
            for sheet, records in _fixture_tables().items():
                header, rows = _rows(records)
                worksheet = book[sheet]
                assert (worksheet.max_row, worksheet.max_column) == (
                    len(rows) + 1,
                    len(header),
                ), sheet
        finally:
            book.close()


# ── Same DataFrames as the upstream xlsx ───────────────────────────────────


def _without_empty_columns(frame):
    """Drop the columns that hold no value at all, which the JSON cannot carry
    (TestSurvivorSheetRows). The fixtures' trimming leaves a few such columns,
    for instance Vote History's split_vote; none occur in the full tables."""
    return frame.loc[:, frame.notna().any()]


def _in_a_fixed_order(frame):
    """Rows sorted on every column, so two exports that order rows differently
    can be compared row by row. Floats are rounded to the JSON's 4 decimals
    for the sort only, so a rounded value sorts where the original did."""
    key = frame.copy()
    for col in key.columns:
        if pd.api.types.is_float_dtype(key[col]):
            key[col] = key[col].round(4)
    order = key.astype(str).sort_values(sorted(key.columns), kind="stable").index
    return frame.loc[order].reset_index(drop=True)


def _sheet_names():
    from app.data import SURVIVOR_TABLES

    return list(SURVIVOR_TABLES)


class TestSameDataFramesAsTheUpstreamXlsx:
    """Each sheet, read with pd.read_excel as the app reads it, from the
    upstream xlsx and from the workbook built out of the same commit's JSON.

    Row order and column order are not compared: survivoR's two exports list
    rows in different orders, and the app only reads columns by name. That
    nothing the app computes depends on row order, or on the empty columns
    the JSON cannot carry, is checked below by running refresh_season on
    both files.
    """

    @pytest.mark.parametrize("sheet", _sheet_names())
    def test_same_columns_and_dtypes(self, sheet, built_xlsx):
        upstream = _without_empty_columns(pd.read_excel(FIXTURE_XLSX, sheet))
        built = pd.read_excel(built_xlsx, sheet)

        assert sorted(built.columns) == sorted(upstream.columns)
        assert built.dtypes[upstream.columns].astype(str).to_dict() == (
            upstream.dtypes.astype(str).to_dict()
        )

    @pytest.mark.parametrize("sheet", _sheet_names())
    def test_same_values(self, sheet, built_xlsx):
        upstream = _in_a_fixed_order(
            _without_empty_columns(pd.read_excel(FIXTURE_XLSX, sheet))
        )
        built = _in_a_fixed_order(pd.read_excel(built_xlsx, sheet))

        assert len(built) == len(upstream) > 0
        for col in sorted(upstream.columns):
            if pd.api.types.is_float_dtype(upstream[col]):
                pd.testing.assert_series_equal(
                    built[col],
                    upstream[col],
                    check_exact=False,
                    rtol=0,
                    atol=JSON_DECIMALS_TOLERANCE,
                )
            else:
                pd.testing.assert_series_equal(built[col], upstream[col])

    def test_dates_bools_and_gaps_survive_the_round_trip(self, built_xlsx):
        """The three shapes most likely to drift, named so a failure says which."""
        episodes = pd.read_excel(built_xlsx, "Episodes")
        castaways = pd.read_excel(built_xlsx, "Castaways")
        summary = pd.read_excel(built_xlsx, "Season Summary")

        assert pd.api.types.is_datetime64_any_dtype(episodes["episode_date"])
        # A bool column with a gap reads as 1.0 / 0.0 / NaN, as upstream's does
        assert castaways["finalist"].dtype == "float64"
        assert castaways["finalist"].isna().any()
        # An in-progress season does not know its finalists yet
        s51 = summary[summary["season"] == 51].iloc[0]
        assert pd.isna(s51["n_finalists"])


# ── Same season in the app ─────────────────────────────────────────────────


@pytest.fixture()
def app(tmp_path, monkeypatch):
    """Flask app with a temp file-backed SQLite DB."""
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path / 'test.db'}")
    if "config" in sys.modules:
        importlib.reload(sys.modules["config"])
    monkeypatch.setattr("app.scheduler.init_scheduler", lambda _app: None)

    from app import create_app, db

    application = create_app()
    application.config["TESTING"] = True
    with application.app_context():
        yield db
        db.session.remove()

    monkeypatch.delenv("DATABASE_URL", raising=False)
    if "config" in sys.modules:
        importlib.reload(sys.modules["config"])


def _refresh_from(db, monkeypatch, xlsx, number):
    """Run refresh_season for a fresh season from one xlsx; return what it stored."""
    from app import data
    from app.models import Season, Survivor

    monkeypatch.setattr(data, "SURVIVOR_DATA_FILE", str(xlsx))
    season = Season(number=number, name=f"Season {number}", is_active=False)
    db.session.add(season)
    db.session.commit()
    data.refresh_season(season)

    skip = {"id", "season_id"}
    stored = {
        "season": {
            c.name: getattr(season, c.name)
            for c in Season.__table__.columns
            if c.name != "id"
        }
    }
    for surv in Survivor.query.filter_by(season_id=season.id):
        row = {
            c.name: getattr(surv, c.name)
            for c in Survivor.__table__.columns
            if c.name not in skip
        }
        row["episode_stats"] = json.loads(row["episode_stats"] or "null")
        stored[surv.castaway_id] = row

    Survivor.query.filter_by(season_id=season.id).delete()
    db.session.delete(season)
    db.session.commit()
    return stored


class TestSameSeasonInTheApp:
    @pytest.mark.parametrize("number", [50, 51])
    def test_refresh_season_stores_the_same_season(
        self, number, app, monkeypatch, built_xlsx
    ):
        upstream = _refresh_from(app, monkeypatch, FIXTURE_XLSX, number)
        built = _refresh_from(app, monkeypatch, built_xlsx, number)

        assert set(built) == set(upstream)
        assert len(upstream) > 1
        for key, fields in upstream.items():
            if key == "season":
                assert built[key] == fields
                continue
            score = fields.pop("performance_score")
            built_score = built[key].pop("performance_score")
            assert built[key] == fields, key
            if score is None:
                assert built_score is None
            else:
                assert built_score == pytest.approx(score, abs=JSON_DECIMALS_TOLERANCE)


class TestOneWorkbookPerRefresh:
    """download_survivor_data() can swap a new build in at any moment.
    refresh_season() read each sheet by path, so one refresh could read some
    sheets from the old build and the rest from the new. It now opens the
    file once and reads every sheet from that, and an open file keeps
    reading the build it opened."""

    def test_a_build_swapped_in_mid_refresh_is_not_read(
        self, app, monkeypatch, built_xlsx, tmp_path
    ):
        import os
        import shutil

        from app import data

        xlsx = tmp_path / "survivoR.xlsx"
        shutil.copy(built_xlsx, xlsx)
        undisturbed = _refresh_from(app, monkeypatch, xlsx, 51)
        real_read_excel = data.pd.read_excel
        reads = []

        def read_then_swap(source, *args, **kwargs):
            frame = real_read_excel(source, *args, **kwargs)
            reads.append(args[:1])
            if len(reads) == 1:
                newer = tmp_path / "newer.xlsx"
                newer.write_bytes(b"a build that is not the one it opened")
                os.replace(newer, xlsx)
            return frame

        monkeypatch.setattr(data.pd, "read_excel", read_then_swap)

        swapped = _refresh_from(app, monkeypatch, xlsx, 51)

        assert len(reads) > 1
        assert swapped == undisturbed
